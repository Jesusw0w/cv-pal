"""Writing the tailored CV as a `.docx`.

Markdown is the content; **this is the submission format**. Applicant tracking systems
ingest `.docx` and `.pdf`, and `.docx` is the one worth doing first because it is a
structured format the parser reads rather than a page description it has to reconstruct.
`python-docx` is already a dependency, used for reading uploads.

## What "ATS-safe" means concretely, and why each rule is here

Every rule below removes a shape this project's *own* `analysis.parseability` reports as
lost — which is the point: the generator and the checker have to agree, or the app flags
documents it produced itself.

- **Built-in heading styles**, not bold body text. A parser looks at the style name to
  decide a line is a section heading; visually-bold Normal text carries no such signal.
- **Nothing in the header or footer.** Contact details there are the classic way an
  email disappears: many parsers read the body story only. So contact goes in the body,
  first.
- **No tables, no text boxes, no columns.** These reorder or drop text on extraction,
  and multi-column layout is exactly what the short-line-ratio check detects.
- **No images.** They carry no text and some parsers reject documents containing them.
- **One paragraph per block, in reading order.** The extraction order is the document
  order, so there is nothing for a parser to get wrong.

The document is deliberately plain. A visually striking CV that an ATS drops is worth
less than a dull one it reads, and the user can restyle the output if a human is the
first reader.
"""

import io
from collections.abc import Sequence

from docx import Document
from docx.shared import Pt

from cv_pal.generation.tailored_cv import Block, BlockKind

# `0` is Word's Title style. Two levels because a CV has two — section, then entry.
_HEADING_LEVEL: dict[BlockKind, int] = {
    BlockKind.NAME: 0,
    BlockKind.SECTION: 1,
    BlockKind.ENTRY: 2,
}


def render_docx(blocks: Sequence[Block]) -> bytes:
    """Render the document as a `.docx` file.

    Args:
        blocks: The document, from `tailored_cv`.

    Returns:
        The file's bytes.
    """
    document = Document()

    # Word's default varies by version and locale, and a substituted font repaginates
    # a document the user has already checked.
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    for block in blocks:
        level = _HEADING_LEVEL.get(block.kind)
        if level is not None:
            document.add_heading(block.text, level=level)
            continue

        paragraph = document.add_paragraph(block.text)
        if block.kind in {BlockKind.HEADLINE, BlockKind.CONTACT, BlockKind.META}:
            # Real paragraphs in the main story — never a header, never a text box.
            for run in paragraph.runs:
                run.font.size = Pt(10)

    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


__all__ = ["render_docx"]
